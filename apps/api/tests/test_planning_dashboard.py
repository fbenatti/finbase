import io
from types import SimpleNamespace

import openpyxl
import pytest
from fastapi import HTTPException

from app.models.planning import BudgetDefenseRequest
from app.routers import planning
from app.routers.planning import (
    _defense_request_out,
    _scope_rows,
    _subtree_ids,
    _timeline_event_matches_budget_scope,
)


def test_global_budget_snapshot_is_visible_inside_department_scope():
    event = {"scope_id_orcamento": None}

    assert _timeline_event_matches_budget_scope(event, {"EXTERNOS"}) is True


def test_budget_snapshot_from_another_scope_is_hidden():
    event = {"scope_id_orcamento": "ID2532-GCP"}

    assert _timeline_event_matches_budget_scope(event, {"EXTERNOS"}) is False


def test_matching_budget_snapshot_is_visible():
    event = {"scope_id_orcamento": "EXTERNOS"}

    assert _timeline_event_matches_budget_scope(event, {"EXTERNOS"}) is True


@pytest.mark.asyncio
async def test_legacy_unscoped_snapshot_rows_are_filtered_by_selected_cloud(monkeypatch):
    departments = {
        1: SimpleNamespace(id=1, parent_id=None, name="Consumo Cloud", id_orcamento=None, active=True),
        2: SimpleNamespace(id=2, parent_id=1, name="Showback", id_orcamento=None, active=True),
        3: SimpleNamespace(id=3, parent_id=2, name="Azure", id_orcamento="EXTERNOS"),
        4: SimpleNamespace(id=4, parent_id=2, name="GCP", id_orcamento="EXTERNOS"),
        5: SimpleNamespace(id=5, parent_id=3, name="Antifraude", id_orcamento=None),
        6: SimpleNamespace(id=6, parent_id=4, name="Outro projeto", id_orcamento=None),
    }

    async def organization(_db):
        return departments, {}

    monkeypatch.setattr(planning, "_organization", organization)
    rows = [
        {"department_id": 5, "values": [10_000.0] * 12},
        {"department_id": 6, "values": [20_000.0] * 12},
    ]

    scoped = await _scope_rows(rows, "EXTERNOS", "Azure", db=None)

    assert scoped == [rows[0]]


@pytest.mark.asyncio
async def test_parse_excel_ignores_empty_trailing_header_cells(monkeypatch):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append([*planning._headers(7), "", ""])
    content = io.BytesIO()
    workbook.save(content)

    async def organization(_db):
        return {}, {}

    monkeypatch.setattr(planning, "_organization", organization)
    assert await planning._parse_excel(content.getvalue(), db=None) == []


@pytest.mark.asyncio
async def test_parse_excel_uses_the_sheet_with_planning_headers(monkeypatch):
    workbook = openpyxl.Workbook()
    workbook.active.append(["Cloud", "Valor Realizado", "Provisão"])
    planning_sheet = workbook.create_sheet("Importação")
    planning_sheet.append(planning._headers(7))
    content = io.BytesIO()
    workbook.save(content)

    async def organization(_db):
        return {}, {}

    monkeypatch.setattr(planning, "_organization", organization)
    assert await planning._parse_excel(content.getvalue(), db=None) == []


@pytest.mark.asyncio
async def test_parse_excel_ignores_blank_department_levels(monkeypatch):
    root = SimpleNamespace(id=1, parent_id=None, name="Raiz", active=True)
    child = SimpleNamespace(id=2, parent_id=1, name="Filho", active=True)
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append(planning._headers(7))
    worksheet.append(["Raiz", "", "Filho", "", "", "", "", *([0] * 12)])
    content = io.BytesIO()
    workbook.save(content)

    async def organization(_db):
        return {1: root, 2: child}, {None: [root], 1: [child], 2: []}

    monkeypatch.setattr(planning, "_organization", organization)
    rows = await planning._parse_excel(content.getvalue(), db=None)

    assert [row.department_id for row in rows] == [2]


@pytest.mark.asyncio
async def test_parse_excel_reports_received_headers_for_detailed_workbook():
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    received = [*planning._level_headers(7), "FOCUS", "Valor de referência", *planning.MONTHS]
    worksheet.append(received)
    content = io.BytesIO()
    workbook.save(content)

    with pytest.raises(HTTPException) as error:
        await planning._parse_excel(content.getvalue(), db=None)

    assert error.value.detail["message"] == "Modelo Excel inválido"
    assert ", ".join(received) in error.value.detail["errors"][1]


@pytest.mark.asyncio
async def test_individual_line_is_kept_even_outside_the_inherited_scope(monkeypatch):
    """A project defended under another ID Orçamento appears as an individual line."""
    departments = {
        1: SimpleNamespace(id=1, parent_id=None, name="Consumo Cloud", id_orcamento=None, active=True),
        2: SimpleNamespace(id=2, parent_id=1, name="Showback", id_orcamento=None, active=True),
        3: SimpleNamespace(id=3, parent_id=2, name="Azure", id_orcamento="ID2531-AZURE", active=True),
        4: SimpleNamespace(id=4, parent_id=3, name="Projeto Existente", id_orcamento="EXTERNOS", active=True),
    }

    async def organization(_db):
        return departments, {}

    monkeypatch.setattr(planning, "_organization", organization)
    rows = [
        {"department_id": 4, "values": [100.0] * 12},
        {"department_id": 3, "values": [0.0] * 12},
    ]

    without_allowlist = await _scope_rows(rows, "ID2531-AZURE", "Azure", db=None)
    assert without_allowlist == [rows[1]]

    with_allowlist = await _scope_rows(rows, "ID2531-AZURE", "Azure", db=None, allow_department_ids={4})
    assert with_allowlist == [rows[0], rows[1]]


def test_subtree_ids_walks_every_descendant():
    departments = {
        1: SimpleNamespace(id=1, parent_id=None),
        2: SimpleNamespace(id=2, parent_id=1),
        3: SimpleNamespace(id=3, parent_id=1),
        4: SimpleNamespace(id=4, parent_id=2),
    }
    children = {}
    for item in departments.values():
        children.setdefault(item.parent_id, []).append(item)

    assert _subtree_ids(1, children) == {1, 2, 3, 4}
    assert _subtree_ids(2, children) == {2, 4}
    assert _subtree_ids(4, children) == {4}


def test_new_project_is_not_an_individual_line():
    departments = {
        1: SimpleNamespace(id=1, parent_id=None, name="Consumo Cloud", id_orcamento=None, active=True),
        2: SimpleNamespace(id=2, parent_id=1, name="Showback", id_orcamento=None, active=True),
        3: SimpleNamespace(id=3, parent_id=2, name="Azure", id_orcamento="ID2531-AZURE", active=True),
        4: SimpleNamespace(id=4, parent_id=3, name="Projeto Novo", id_orcamento="NOVO-ID", active=True),
    }
    request = BudgetDefenseRequest(
        year=2026,
        department_id=4,
        requested_id_orcamento="NOVO-ID",
        requested_cloud="Azure",
        status="SOLICITADA",
        is_new_project=True,
    )

    result = _defense_request_out(request, departments)

    assert result["is_individual_line"] is False
    assert result["project_name"] == "Projeto Novo"
    assert result["organograma_id_orcamento"] == "NOVO-ID"


def test_existing_project_with_another_requested_id_is_an_individual_line():
    departments = {
        1: SimpleNamespace(id=1, parent_id=None, name="Consumo Cloud", id_orcamento=None, active=True),
        2: SimpleNamespace(id=2, parent_id=1, name="Showback", id_orcamento=None, active=True),
        3: SimpleNamespace(id=3, parent_id=2, name="Azure", id_orcamento="ID2531-AZURE", active=True),
        4: SimpleNamespace(id=4, parent_id=3, name="Projeto Existente", id_orcamento="EXTERNOS", active=True),
    }
    request = BudgetDefenseRequest(
        year=2026,
        department_id=4,
        requested_id_orcamento="ID2531-AZURE",
        requested_cloud="Azure",
        status="SOLICITADA",
    )

    result = _defense_request_out(request, departments)

    assert result["is_individual_line"] is True
    # The organogram keeps its own line until the defense is published.
    assert result["organograma_id_orcamento"] == "EXTERNOS"

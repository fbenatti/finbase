from __future__ import annotations

import io
import json
import re
from collections import defaultdict
from datetime import date, datetime, timezone
from typing import Any

import openpyxl
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.deps import get_current_user, require
from app.database import get_db
from app.models.destination import DestinationTable
from app.models.identity import User
from app.models.planning import BudgetDefenseRequest, FinancialType, PlanningEntry, PlanningVersion
from app.models.contract_story import ContractStoryLedger
from app.models.budget_accuracy import BudgetAccuracyAck
from app.models.rules import Department
from app.services.trino_service import execute as trino_execute
from app.services import lakehouse_layers as layers
from app.services.gold_materialization import GOLD

router = APIRouter(prefix="/planning", tags=["planning"])

MONTHS = ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"]
LEVELS = [f"Departamento nível {i}" for i in range(1, 8)]
HEADERS = [*LEVELS, *MONTHS]
SYSTEM_CODES = {"realizado", "forecast"}
DEFAULT_TYPES = [("realizado", "Realizado"), ("forecast", "Forecast (projetado)"), ("saldo-contratual", "Saldo contratual"), ("orcamento", "Orçamento"), ("forecast-manual", "Forecast manual"), ("base-line-contratual", "Base Line Contratual"), ("carta-credito", "Carta de Credito"), ("defesa-orcamentaria", "Defesa Orçamentária")]
EVENT_TYPES = {"CARGA_INICIAL", "DEFESA", "APROVACAO", "CORTE", "AJUSTE", "INCLUSAO_PROJETO", "REALOCACAO", "RESTAURACAO", "OUTRO"}

# Defesa orçamentária de projeto.  A solicitação vive entre o organograma e o
# orçamento: enquanto aberta, o projeto defendido pode ocupar uma "linha
# individual" no ID Orçamento pedido sem alterar o organograma; ao publicar, o
# ID final é gravado no nó, garantindo que o projeto pertença a uma única linha.
DEFENSE_OPEN_STATUSES = ("RASCUNHO", "SOLICITADA", "APROVADA")
DEFENSE_STATUSES = ("RASCUNHO", "SOLICITADA", "APROVADA", "PUBLICADA", "REJEITADA")
DEFENSE_SCOPE_REQUEST = "defesa-solicitar"
DEFENSE_SCOPE_PUBLISH = "defesa-publicar"


class PlanningRowIn(BaseModel):
    entity_type: str
    department_id: int
    values: list[float] = Field(min_length=12, max_length=12)


class PlanningSaveIn(BaseModel):
    rows: list[PlanningRowIn]
    title: str
    description: str
    event_type: str
    effective_date: date
    scope_id_orcamento: str | None = None
    scope_cloud: str | None = None
    # Quando informado, a defesa pertence a uma solicitação: o escopo vem dela e
    # a versão vigente é controlada por solicitação, não por (ID, cloud).
    defense_request_id: int | None = None
    submit: bool = False


class PlanningRestoreIn(BaseModel):
    title: str
    description: str
    event_type: str = "RESTAURACAO"
    effective_date: date
    scope_id_orcamento: str | None = None
    scope_cloud: str | None = None


class BudgetPublishIn(BaseModel):
    id_orcamento: str
    cloud_scope: str | None = None
    title: str
    description: str
    effective_date: date


class ContractStoryLedgerIn(BaseModel):
    contracts: list[dict[str, Any]] = Field(default_factory=list)
    months: dict[str, list[dict[str, Any]]] = Field(default_factory=dict)
    milestones: dict[str, list[dict[str, Any]]] = Field(default_factory=dict)


def _merge_ledger_rows(existing: list[dict[str, Any]], incoming: list[dict[str, Any]], key: str) -> list[dict[str, Any]]:
    merged = {str(row[key]): row for row in existing if row.get(key) is not None}
    merged.update({str(row[key]): row for row in incoming if row.get(key) is not None})
    return list(merged.values())


@router.get("/contract-story")
async def get_contract_story_ledger(
    db: AsyncSession = Depends(get_db),
    _user: User = Depends(require("planning", "read")),
):
    ledger = await db.get(ContractStoryLedger, 1)
    return ledger.payload if ledger else None


@router.put("/contract-story")
async def save_contract_story_ledger(
    body: ContractStoryLedgerIn,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require("planning", "update")),
):
    ledger = await db.get(ContractStoryLedger, 1)
    payload = {
        "contracts": body.contracts,
        "months": body.months,
        "milestones": body.milestones,
    }
    if ledger:
        current = ledger.payload or {}
        payload["contracts"] = _merge_ledger_rows(current.get("contracts", []), body.contracts, "id")
        payload["months"] = {
            contract_id: _merge_ledger_rows(
                current.get("months", {}).get(contract_id, []),
                body.months.get(contract_id, []),
                "month",
            )
            for contract_id in set(current.get("months", {})) | set(body.months)
        }
        payload["milestones"] = {
            contract_id: _merge_ledger_rows(
                current.get("milestones", {}).get(contract_id, []),
                body.milestones.get(contract_id, []),
                "id",
            )
            for contract_id in set(current.get("milestones", {})) | set(body.milestones)
        }
    else:
        ledger = ContractStoryLedger(id=1)
        db.add(ledger)
    ledger.payload = payload
    ledger.updated_by = user.id
    await db.commit()
    return ledger.payload


class DefenseProjectIn(BaseModel):
    """Novo projeto criado dentro da Defesa Orçamentária."""

    year: int
    parent_id: int
    name: str
    id_orcamento: str
    centro_custo: str | None = None
    ambiente: str | None = None
    cloud_filter: str | None = None
    title: str | None = None
    description: str | None = None


class DefenseRequestIn(BaseModel):
    """Solicitação de defesa de um projeto já existente no organograma."""

    year: int
    department_id: int
    requested_id_orcamento: str
    requested_cloud: str | None = None
    title: str | None = None
    description: str | None = None


class DefenseRequestUpdateIn(BaseModel):
    requested_id_orcamento: str | None = None
    requested_cloud: str | None = None
    title: str | None = None
    description: str | None = None


class DefenseApproveIn(BaseModel):
    final_id_orcamento: str
    final_cloud: str | None = None
    decision_note: str
    effective_date: date | None = None


class DefensePublishIn(BaseModel):
    title: str
    description: str
    effective_date: date
    final_id_orcamento: str | None = None
    final_cloud: str | None = None
    decision_note: str | None = None


class DefenseRejectIn(BaseModel):
    decision_note: str


class BudgetAccuracyAckIn(BaseModel):
    target_month: str = Field(min_length=6, max_length=6)
    planning_version_id: int | None = None
    department_id: int | None = None
    department: str
    budget_value: float = 0
    actual_value: float = 0
    category: str
    reason: str
    severity: str = "Informação"
    ordem: int = 0
    registered_at: date | None = None


def _out(row: Any) -> dict[str, Any]:
    return {c.name: getattr(row, c.name) for c in row.__table__.columns}


async def ensure_default_financial_types(db: AsyncSession) -> None:
    existing = set((await db.execute(select(FinancialType.code))).scalars().all())
    for code, name in DEFAULT_TYPES:
        if code not in existing:
            db.add(FinancialType(code=code, name=name, active=True))
    await db.commit()


async def ensure_default_budget_ids(db: AsyncSession) -> int:
    """Seed budget IDs on Cloud nodes without replacing user-maintained values."""
    departments = (await db.execute(select(Department))).scalars().all()
    by_id = {item.id: item for item in departments}
    showback = {"azure": "ID2531-AZURE", "gcp": "ID2532-GCP", "aws": "ID2533-AWS", "oci": "ID2534-OCI"}
    updated = 0
    for department in departments:
        parent = by_id.get(department.parent_id)
        if not parent or department.id_orcamento:
            continue
        cloud = (department.name or "").strip().lower()
        if parent.name.strip().lower() == "showback" and cloud in showback:
            department.id_orcamento = showback[cloud]
            updated += 1
        elif parent.name.strip().lower() == "chargeback" and cloud:
            department.id_orcamento = f"EXTERNO-{cloud.upper()}"
            updated += 1
    if updated:
        await db.commit()
    return updated


async def ensure_planning_paths_synced(db: AsyncSession) -> int:
    departments = {d.id: d for d in (await db.execute(select(Department))).scalars().all()}
    entries = (await db.execute(select(PlanningEntry).where(PlanningEntry.department_id != None))).scalars().all()
    updated = 0
    for entry in entries:
        dept = departments.get(entry.department_id)
        if dept:
            path: list[str] = []
            cur = dept
            seen = set()
            while cur and cur.id not in seen:
                seen.add(cur.id)
                path.append((cur.name or "").strip())
                cur = departments.get(cur.parent_id) if cur.parent_id else None
            path.reverse()
            if entry.department_levels != path:
                entry.department_levels = path
                updated += 1
    if updated:
        await db.commit()
    return updated


async def _type_for_write(code: str, db: AsyncSession) -> FinancialType:
    row = (await db.execute(select(FinancialType).where(FinancialType.code == code))).scalar_one_or_none()
    if not row:
        raise HTTPException(404, "Tipificação financeira não encontrada")
    if not row.active:
        raise HTTPException(422, "Tipificação financeira inativa")
    if row.code in SYSTEM_CODES:
        raise HTTPException(422, "Realizado e forecast automático não aceitam carga manual")
    return row


async def _organization(db: AsyncSession) -> tuple[dict[int, Department], dict[int | None, list[Department]]]:
    departments = (await db.execute(select(Department))).scalars().all()
    by_id = {d.id: d for d in departments}
    children: dict[int | None, list[Department]] = defaultdict(list)
    for d in departments:
        children[d.parent_id].append(d)
    return by_id, children


def _department_scope(department: Department, departments: dict[int, Department]) -> tuple[str | None, str | None]:
    chain: list[Department] = []
    current: Department | None = department
    while current:
        chain.append(current)
        current = departments.get(current.parent_id) if current.parent_id else None
    chain.reverse()
    budget_id = next((item.id_orcamento for item in reversed(chain) if item.id_orcamento), None)
    cloud = chain[2].name if len(chain) >= 3 else None
    return budget_id, cloud


async def _scope_rows(rows: list[dict[str, Any]], budget_id: str | None, cloud: str | None, db: AsyncSession, allow_department_ids: set[int] | None = None) -> list[dict[str, Any]]:
    """Restrict rows to the selected budget scope.

    ``allow_department_ids`` keeps departments that live outside the inherited
    scope but are legitimately part of it — an *individual line* from an open
    defense request, or the project being defended.
    """
    if not budget_id and not cloud:
        return rows
    departments, _ = await _organization(db)
    result: list[dict[str, Any]] = []
    for row in rows:
        department = departments.get(row.get("department_id"))
        if department is None:
            continue
        row_budget, row_cloud = _department_scope(department, departments)
        in_scope = (not budget_id or row_budget == budget_id) and (not cloud or row_cloud == cloud)
        if in_scope or (allow_department_ids and department.id in allow_department_ids):
            result.append(row)
    return result


def _subtree_ids(root_id: int, children: dict[int | None, list[Department]]) -> set[int]:
    result: set[int] = set()
    pending = [root_id]
    while pending:
        current = pending.pop()
        if current in result:
            continue
        result.add(current)
        pending.extend(child.id for child in children.get(current, []))
    return result


def _defense_request_out(request: BudgetDefenseRequest, departments: dict[int, Department]) -> dict[str, Any]:
    department = departments.get(request.department_id)
    org_id, org_cloud = _department_scope(department, departments) if department else (None, None)
    # A new project is created already carrying its own ID Orçamento, so it is
    # not an individual line.  An existing project defended under a different ID
    # Orçamento is: the organograma still points at another line until publish.
    is_individual_line = bool(department) and not (
        org_id == request.requested_id_orcamento
        and (org_cloud or None) == (request.requested_cloud or None)
    )
    return {
        **_out(request),
        "project_name": department.name if department else None,
        "project_active": bool(department.active) if department else False,
        "project_path": _levels(department, departments) if department else [],
        "organograma_id_orcamento": org_id,
        "organograma_cloud": org_cloud,
        "is_individual_line": is_individual_line,
    }


async def _load_defense_request(db: AsyncSession, request_id: int, year: int | None = None) -> BudgetDefenseRequest:
    request = await db.get(BudgetDefenseRequest, request_id)
    if not request or (year is not None and request.year != year):
        raise HTTPException(404, "Solicitação de defesa orçamentária não encontrada")
    return request


async def _defense_requests_for_scope(db: AsyncSession, year: int, budget_id: str | None, cloud: str | None, statuses: tuple[str, ...] = DEFENSE_OPEN_STATUSES) -> list[BudgetDefenseRequest]:
    if not budget_id:
        return []
    query = select(BudgetDefenseRequest).where(
        BudgetDefenseRequest.year == year,
        BudgetDefenseRequest.requested_id_orcamento == budget_id,
        BudgetDefenseRequest.status.in_(statuses),
    ).order_by(BudgetDefenseRequest.id)
    rows = (await db.execute(query)).scalars().all()
    if cloud:
        # Requests created before cloud scoping keep requested_cloud null and
        # still belong to the selected cloud.
        rows = [row for row in rows if not row.requested_cloud or row.requested_cloud == cloud]
    return list(rows)


async def _individual_line_map(db: AsyncSession, year: int, budget_id: str | None, cloud: str | None, departments: dict[int, Department]) -> dict[int, BudgetDefenseRequest]:
    """department_id -> open request for projects defended outside their line.

    These are the "individual lines": the project is defended under the
    requested ID Orçamento while the organograma keeps its current one.
    """
    result: dict[int, BudgetDefenseRequest] = {}
    for request in await _defense_requests_for_scope(db, year, budget_id, cloud):
        department = departments.get(request.department_id)
        if not department:
            continue
        org_id, org_cloud = _department_scope(department, departments)
        if org_id == request.requested_id_orcamento and (org_cloud or None) == (request.requested_cloud or None):
            continue
        result[department.id] = request
    return result


async def _ensure_defense_open(year: int, budget_id: str, cloud: str | None, db: AsyncSession) -> None:
    defense_type_id = (await db.execute(select(FinancialType.id).where(FinancialType.code == "defesa-orcamentaria"))).scalar_one_or_none()
    approved = (await db.execute(select(PlanningVersion.id).where(PlanningVersion.financial_type_id == defense_type_id, PlanningVersion.year == year, PlanningVersion.scope_id_orcamento == budget_id, PlanningVersion.scope_cloud == cloud, PlanningVersion.published_to_budget_at != None).limit(1))).scalar_one_or_none()
    if approved:
        raise HTTPException(422, "A defesa deste ID Orçamento/Cloud já foi aprovada; publique a alteração como revisão no módulo Orçamento")


def _cloud_origem_sql(columns: dict[str, str]) -> str | None:
    column = columns.get("cloud_origem")
    return f'LOWER(CAST("{column}" AS VARCHAR))' if column else None


def _cloud_filter_sql(cloud: str | None, columns: dict[str, str]) -> str:
    if not cloud:
        return ""
    slug = cloud.strip().lower()
    if not re.fullmatch(r"[a-z0-9]+(?:-[a-z0-9]+)*", slug):
        raise HTTPException(422, "Filtro de cloud inválido")
    expr = _cloud_origem_sql(columns)
    if not expr:
        raise HTTPException(422, "A tabela Gold não possui cloud_origem")
    if slug in {"gcp", "azure", "atlas"}:
        return f"AND {expr} LIKE '{slug}%'"
    return f"AND {expr} = '{slug}'"


def _levels(department: Department, departments: dict[int, Department]) -> list[str]:
    chain: list[str] = []
    current: Department | None = department
    while current:
        chain.append(current.name)
        current = departments.get(current.parent_id) if current.parent_id else None
    return list(reversed(chain))


def _budget_accuracy_metrics(budget: float, actual: float) -> dict[str, Any]:
    variance = actual - budget
    return {
        "budget_value": budget,
        "actual_value": actual,
        "variance": variance,
        "variance_pct": variance / budget if budget else None,
        "consumed_pct": actual / budget if budget else None,
        "status": "sem_orcamento" if not budget and not actual else "sem_realizado" if not actual else "fora_orcamento" if variance > 0 else "dentro_orcamento",
    }


async def _budget_accuracy_actuals(
    year: int,
    months: list[int],
    scope_cloud: str | None = None,
    scope_path: list[str] | None = None,
) -> list[dict[str, Any]]:
    levels = ", ".join(
        f'CAST("organograma_nivel_{level}" AS VARCHAR) AS organograma_nivel_{level}'
        for level in range(1, 8)
    )
    periods = ", ".join(f"'{year}{month:02d}'" for month in months)
    cloud_filter = ""
    if scope_cloud:
        cloud = scope_cloud.strip().lower()
        if cloud == "azure":
            cloud_filter = " AND lower(CAST(\"cloud_origem\" AS VARCHAR)) LIKE 'azure%'"
        elif cloud == "gcp":
            cloud_filter = " AND lower(CAST(\"cloud_origem\" AS VARCHAR)) LIKE 'gcp%'"
        else:
            cloud_filter = f" AND lower(CAST(\"cloud_origem\" AS VARCHAR)) = '{cloud}'"
    scope_filter = ""
    if scope_path:
        escaped = [item.strip().replace("'", "''") for item in scope_path]
        candidates = [list(range(1, len(escaped) + 1))]
        if len(escaped) >= 3:
            candidates.append([1, 2, 4, 3, *range(5, len(escaped) + 2)])
        # A project under Showback and the same project under Chargeback are
        # one consumption entity for budget accuracy. Include both branches
        # and support both historical Azure/Diretoria level orders.
        if len(escaped) >= 5 and escaped[1] in {"Showback", "Chargeback"}:
            candidates.extend([
                [1, 3, 4, *range(5, len(escaped) + 2)],
                [1, 4, 3, *range(5, len(escaped) + 2)],
            ])
        candidate_sql = []
        for levels_order in candidates:
            checks = [
                f'TRIM(CAST("organograma_nivel_{physical_level}" AS VARCHAR)) = \'{escaped[index]}\''
                for index, physical_level in enumerate(levels_order[:len(escaped)])
            ]
            candidate_sql.append(" AND ".join(checks))
        scope_filter = " AND (" + " OR ".join(f"({item})" for item in candidate_sql) + ")"
    group_levels = ", ".join(f'"organograma_nivel_{level}"' for level in range(1, 8))
    sql = f"""
        SELECT {levels}, period, SUM(actual_value) AS actual_value, SUM(forecast_value) AS forecast_value, MAX(has_actual) AS has_actual
        FROM (
          SELECT {levels}, CAST("_billing_period" AS VARCHAR) AS period,
                 SUM(COALESCE(TRY_CAST("realizado" AS DOUBLE), 0)) AS actual_value,
                 CAST(0 AS DOUBLE) AS forecast_value, 1 AS has_actual
          FROM iceberg.{GOLD}.dest_dm_cost_visao_completa
          WHERE CAST("_billing_period" AS VARCHAR) IN ({periods})
            AND ("tipo" IS NULL OR TRIM(CAST("tipo" AS VARCHAR)) IN ('', 'realizado')){cloud_filter}
          {scope_filter}
          GROUP BY {group_levels}, "_billing_period"
          UNION ALL
          SELECT {levels}, CAST("period" AS VARCHAR) AS period,
                 CAST(0 AS DOUBLE) AS actual_value,
                 SUM(COALESCE(TRY_CAST("realizado" AS DOUBLE), 0)) AS forecast_value, 0 AS has_actual
          FROM iceberg.{GOLD}.dest_dm_cost_visao_completa
          WHERE CAST("period" AS VARCHAR) IN ({periods})
            AND lower(TRIM(CAST("tipo" AS VARCHAR))) = 'forecast'{cloud_filter}
          {scope_filter}
          GROUP BY {group_levels}, "period"
        ) facts
        GROUP BY {group_levels}, period
    """
    return await trino_execute(sql, timeout_s=180)


async def _validate_rows(rows: list[PlanningRowIn], db: AsyncSession) -> list[dict[str, Any]]:
    departments, _ = await _organization(db)
    prepared: list[dict[str, Any]] = []
    by_entity: dict[tuple[str, int], dict[str, Any]] = {}
    errors: list[str] = []
    for pos, row in enumerate(rows, 1):
        entity_type = row.entity_type.lower()
        if entity_type != "department":
            errors.append(f"Linha {pos}: entity_type deve ser department")
            continue
        department = departments.get(row.department_id)
        if not department or not department.active:
            errors.append(f"Linha {pos}: departamento inválido ou inativo")
            continue
        key = (entity_type, department.id)
        existing = by_entity.get(key)
        if existing:
            existing["values"] = [
                float(current) + float(incoming)
                for current, incoming in zip(existing["values"], row.values)
            ]
            continue
        prepared_row = {"entity_type": entity_type, "department_id": department.id, "department_levels": _levels(department, departments), "values": [float(value) for value in row.values]}
        by_entity[key] = prepared_row
        prepared.append(prepared_row)
    if errors:
        raise HTTPException(422, {"message": "Dados inválidos", "errors": errors})
    return prepared


def _validate_event(title: str, description: str, event_type: str, effective_date: date) -> tuple[str, str, str, date]:
    title, description, event_type = title.strip(), description.strip(), event_type.strip().upper()
    if not title or not description or not effective_date:
        raise HTTPException(422, "Título, descrição, tipo do evento e data efetiva são obrigatórios")
    if event_type not in EVENT_TYPES:
        raise HTTPException(422, "Tipo de evento inválido")
    return title, description, event_type, effective_date


async def _publish_version(fin_type: FinancialType, year: int, rows: list[dict[str, Any]], source: str, db: AsyncSession, *, title: str, description: str, event_type: str, effective_date: date, filename: str | None = None, restored_from: int | None = None, scope_id_orcamento: str | None = None, scope_cloud: str | None = None, defense_request_id: int | None = None) -> PlanningVersion:
    if year <= 0:
        raise HTTPException(422, "Ano inválido")
    title, description, event_type, effective_date = _validate_event(title, description, event_type, effective_date)
    current_query = select(PlanningVersion).where(PlanningVersion.financial_type_id == fin_type.id, PlanningVersion.year == year, PlanningVersion.is_current == True)
    if defense_request_id is not None:
        # Uma solicitação de defesa tem a própria linha vigente: dois projetos
        # podem defender o mesmo ID Orçamento sem sobrescrever um ao outro.
        current_query = current_query.where(PlanningVersion.defense_request_id == defense_request_id)
    elif fin_type.code == "defesa-orcamentaria" and scope_id_orcamento:
        current_query = current_query.where(PlanningVersion.scope_id_orcamento == scope_id_orcamento, PlanningVersion.scope_cloud == scope_cloud)
    current = (await db.execute(current_query)).scalars().all()
    next_version = max((v.version for v in current), default=0)
    all_versions = (await db.execute(select(func.max(PlanningVersion.version)).where(PlanningVersion.financial_type_id == fin_type.id, PlanningVersion.year == year))).scalar_one()
    version = int(all_versions or 0) + 1
    # Orçamento é um livro de lançamentos obrigatório: cada importação é uma
    # carga complementar (inclusive quando negativa, para cortes) e deve
    # continuar vigente para compor o saldo. Os demais tipos preservam o
    # comportamento de snapshot, no qual uma publicação substitui a anterior.
    if fin_type.code != "orcamento":
        for old in current:
            old.is_current = False
    snapshot = PlanningVersion(financial_type_id=fin_type.id, year=year, version=version, is_current=True, source=source, source_filename=filename, restored_from_id=restored_from, classification_version_id=None, title=title, description=description, event_type=event_type, effective_date=effective_date, scope_id_orcamento=scope_id_orcamento, scope_cloud=scope_cloud, defense_request_id=defense_request_id)
    db.add(snapshot)
    await db.flush()
    for row in rows:
        for month, value in enumerate(row["values"], 1):
            db.add(PlanningEntry(planning_version_id=snapshot.id, entity_type=row["entity_type"], department_id=row["department_id"], department_levels=row["department_levels"], month=month, value=float(value)))
    await db.commit()
    await db.refresh(snapshot)
    return snapshot


@router.get("/types")
async def list_types(db: AsyncSession = Depends(get_db)):
    await ensure_default_financial_types(db)
    rows = (await db.execute(select(FinancialType).order_by(FinancialType.name))).scalars().all()
    return [_out(r) for r in rows]


@router.get("/template")
async def download_template(year: int, db: AsyncSession = Depends(get_db)):
    if year <= 0: raise HTTPException(422, "Ano inválido")
    departments, _ = await _organization(db)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Cadastro financeiro"; ws.append(HEADERS)
    for department in sorted((d for d in departments.values() if d.active), key=lambda d: _levels(d, departments)):
        levels = _levels(department, departments)
        ws.append([*levels, *([""] * (7 - len(levels))), *([0] * 12)])
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(14, min(35, max(len(str(c.value or "")) for c in col) + 2))
    content = io.BytesIO(); wb.save(content); content.seek(0)
    return StreamingResponse(content, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="modelo-cadastro-financeiro-{year}.xlsx"'})


@router.get("/orcamento/{year}/simulado")
async def download_simulated_budget(
    year: int,
    base_period: str = "202601",
    detailed: bool = False,
    detail_column: str | None = None,
    cloud: str | None = None,
    db: AsyncSession = Depends(get_db),
):
    """Create a simulated annual forecast from a selected realized month.

    The consolidated file is import-ready.  The detailed file is an optional
    budget-defense view split by FOCUS ServiceName when that field exists.
    Use cloud=gcp to generate only GCP project rows (gcp-4p / gcp-ti).
    """
    if year <= 0 or not re.fullmatch(r"\d{6}", base_period):
        raise HTTPException(422, "Ano ou período-base inválido")
    active_destinations = (await db.execute(select(DestinationTable).where(DestinationTable.active == True))).scalars().all()
    destination = next((item for item in active_destinations if "focus" in f"{item.name} {item.physical_name}".lower()), None)
    destination = destination or await db.get(DestinationTable, 1)
    if not destination or not destination.active:
        raise HTTPException(404, "Tabela Gold de realizado não encontrada")
    safe_table = destination.physical_name.replace("-", "_").replace(".", "_")
    try:
        schema_rows = await trino_execute(f"DESCRIBE iceberg.{layers.detail_layer()}.{safe_table}", timeout_s=30)
        columns = {str(row.get("Column") or row.get("column_name") or "").lower(): str(row.get("Column") or row.get("column_name") or "") for row in schema_rows}
        if "_billing_period" not in columns or "realizado" not in columns:
            raise HTTPException(422, "A tabela Gold selecionada não possui _billing_period e realizado")
        level_columns = [columns.get(f"organograma_nivel_{level}") for level in range(1, 8)]
        level_select = [f'CAST("{column}" AS VARCHAR) AS organograma_nivel_{index}' if column else f'CAST(NULL AS VARCHAR) AS organograma_nivel_{index}' for index, column in enumerate(level_columns, 1)]
        suggested_column = next((column for key, column in columns.items() if key.replace("_", "") in {"servicename", "metercategory"}), None)
        selected_detail_column = columns.get((detail_column or "").lower()) or suggested_column
        detail_select = f', COALESCE(NULLIF(TRIM(CAST("{selected_detail_column}" AS VARCHAR)), \'\'), \'Sem sugestão FOCUS\') AS focus_detail_suggestion' if detailed and selected_detail_column else ", 'Sem sugestão FOCUS' AS focus_detail_suggestion" if detailed else ""
        type_filter = f"AND (tipo = 'realizado' OR tipo IS NULL)" if "tipo" in columns else ""
        cloud_filter = _cloud_filter_sql(cloud, columns)
        group_by = ", ".join(str(index) for index in range(1, 8 + (1 if detailed else 0)))
        source_rows = await trino_execute(f"""
            SELECT {', '.join(level_select)}{detail_select},
              SUM(COALESCE(TRY_CAST(realizado AS DOUBLE), 0)) AS value
            FROM iceberg.{layers.detail_layer()}.{safe_table}
            WHERE CAST(_billing_period AS VARCHAR) = '{base_period}'
              {type_filter}
              {cloud_filter}
            GROUP BY {group_by}
        """, timeout_s=120)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(503, f"Não foi possível ler o realizado de {base_period}: {exc}")

    departments, _ = await _organization(db)
    by_path = {tuple(_levels(d, departments)): d for d in departments.values() if d.active}
    baseline_by_department: dict[int, float] = defaultdict(float)
    for row in source_rows:
        path = tuple(str(row.get(f"organograma_nivel_{level}") or "").strip() for level in range(1, 8))
        path = tuple(level for level in path if level)
        department = by_path.get(path)
        if department:
            baseline_by_department[department.id] += float(row.get("value") or 0)

    wb = openpyxl.Workbook(); ws = wb.active
    if detailed:
        ws.title = "Forecast detalhado"
        detail_label = detail_column or suggested_column or "MeterCategory"
        detailed_headers = [*LEVELS, f"{detail_label} FOCUS (sugestão)", "Valor de referência", *MONTHS]
        ws.append(detailed_headers)
        for row in source_rows:
            levels = [str(row.get(f"organograma_nivel_{level}") or "").strip() for level in range(1, 8)]
            value = float(row.get("value") or 0)
            ws.append([*levels, str(row.get("focus_detail_suggestion") or "Sem sugestão FOCUS"), value, *([value] * 12)])
    else:
        ws.title = "Forecast simulado"
        ws.append(HEADERS)
        for department in sorted((d for d in departments.values() if d.active), key=lambda d: _levels(d, departments)):
            levels = _levels(department, departments)
            monthly_value = baseline_by_department.get(department.id, 0.0)
            if cloud and monthly_value == 0:
                continue
            ws.append([*levels, *([""] * (7 - len(levels))), *([monthly_value] * 12)])
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(14, min(35, max(len(str(c.value or "")) for c in col) + 2))
    content = io.BytesIO(); wb.save(content); content.seek(0)
    export_kind = "detalhado" if detailed else "consolidado"
    cloud_suffix = f"-{cloud.lower()}" if cloud else ""
    return StreamingResponse(
        content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="forecast-simulado-{export_kind}{cloud_suffix}-{year}-base-{base_period}.xlsx"'},
    )


def _cell_str(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _sheet_header(sheet: Any) -> list[str]:
    header = [_cell_str(cell) for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), [])]
    while header and not header[-1]:
        header.pop()
    return header


async def _parse_excel(raw: bytes, db: AsyncSession, create_missing_departments: list[list[str]] | None = None) -> list[PlanningRowIn]:
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    except Exception as exc:
        raise HTTPException(422, f"Arquivo Excel inválido: {exc}")
    sheet = next((candidate for candidate in workbook.worksheets if _sheet_header(candidate) == HEADERS), workbook.active)
    header = _sheet_header(sheet)
    if header != HEADERS:
        received = ", ".join(header) if header else "(vazio)"
        raise HTTPException(422, {"message": "Modelo Excel inválido", "errors": [f"Cabeçalhos esperados: {', '.join(HEADERS)}", f"Cabeçalhos recebidos: {received}"]})
    departments, children = await _organization(db)
    parsed: list[PlanningRowIn] = []
    errors: list[str] = []
    missing: dict[tuple[str, ...], list[int]] = defaultdict(list)
    skipped: list[str] = []
    for line, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        vals = list(values)
        if not any(v not in (None, "") for v in vals): continue
        names = [name for name in (_cell_str(value) for value in vals[:7]) if name]
        if not names[0]: errors.append(f"Linha {line}: Departamento nível 1 é obrigatório"); continue
        path = tuple(names)
        parent: int | None = None; department: Department | None = None; broken = False
        for level, name in enumerate(names, 1):
            matches = [d for d in children[parent] if d.active and d.name == name]
            if len(matches) != 1:
                selected_paths = {tuple(item) for item in (create_missing_departments or [])}
                if create_missing_departments is not None and path not in selected_paths:
                    broken = True
                    break
                if path not in selected_paths:
                    missing[path].append(line); broken = True; break
                created = Department(parent_id=parent, name=name, active=True)
                db.add(created)
                await db.flush()
                departments[created.id] = created
                children.setdefault(parent, []).append(created)
                children.setdefault(created.id, [])
                matches = [created]
            department = matches[0]; parent = department.id
        if broken or not department:
            if path in missing and create_missing_departments is not None:
                skipped.append(f"Linha {line}: {' > '.join(names)}")
            continue
        monthly: list[float] = []
        for month, value in zip(MONTHS, vals[7:19]):
            try:
                if value is None or str(value).strip() == "": raise ValueError()
                monthly.append(float(value))
            except (TypeError, ValueError): errors.append(f"Linha {line}: {month} deve ser numérico")
        if len(monthly) == 12: parsed.append(PlanningRowIn(entity_type="department", department_id=department.id, values=monthly))
    if errors: raise HTTPException(422, {"message": "Importação rejeitada", "errors": errors})
    if missing:
        departments_missing = [{"path": list(path), "lines": lines} for path, lines in missing.items()]
        raise HTTPException(422, {"message": "Departamentos ausentes", "errors": [f"Linha {line}: departamento não cadastrado" for lines in missing.values() for line in lines], "departments_missing": departments_missing})
    return parsed


@router.get("/type-years/{type_code}")
async def list_years(type_code: str, db: AsyncSession = Depends(get_db)):
    fin_type = (await db.execute(select(FinancialType).where(FinancialType.code == type_code))).scalar_one_or_none()
    if not fin_type: raise HTTPException(404, "Tipificação financeira não encontrada")
    rows = (await db.execute(select(PlanningVersion.year).where(PlanningVersion.financial_type_id == fin_type.id).distinct().order_by(PlanningVersion.year.desc()))).scalars().all()
    return rows


@router.get("/scopes/options")
async def planning_scope_options(db: AsyncSession = Depends(get_db)):
    departments, _ = await _organization(db)
    values = {_department_scope(item, departments) for item in departments.values() if item.active}
    return [{"id_orcamento": budget_id, "cloud": cloud} for budget_id, cloud in sorted(values, key=lambda item: ((item[0] or ""), (item[1] or ""))) if budget_id]


@router.post("/{type_code}/{year}/import")
async def import_excel(type_code: str, year: int, file: UploadFile = File(...), title: str = Form(...), description: str = Form(...), event_type: str = Form(...), effective_date: date = Form(...), scope_id_orcamento: str | None = Form(None), scope_cloud: str | None = Form(None), defense_request_id: int | None = Form(None), create_missing_departments: str | None = Form(None), db: AsyncSession = Depends(get_db)):
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xltm")): raise HTTPException(422, "Envie um arquivo .xlsx ou .xltm")
    fin_type = await _type_for_write(type_code, db)
    defense_request: BudgetDefenseRequest | None = None
    if defense_request_id:
        if type_code != "defesa-orcamentaria":
            raise HTTPException(422, "Solicitação de defesa só se aplica à Defesa Orçamentária")
        defense_request = await _load_defense_request(db, defense_request_id, year)
        if defense_request.status not in ("RASCUNHO", "SOLICITADA"):
            raise HTTPException(422, "Esta solicitação não aceita mais alterações de valores")
        scope_id_orcamento = defense_request.requested_id_orcamento
        scope_cloud = defense_request.requested_cloud
    try:
        selected_missing = json.loads(create_missing_departments) if create_missing_departments else None
        if selected_missing is not None and not isinstance(selected_missing, list):
            raise ValueError()
    except (TypeError, ValueError, json.JSONDecodeError):
        raise HTTPException(422, "create_missing_departments deve ser uma lista JSON de caminhos")
    rows = await _parse_excel(await file.read(), db, selected_missing)
    prepared = await _validate_rows(rows, db)
    if type_code == "defesa-orcamentaria" and not scope_id_orcamento:
        raise HTTPException(422, "ID Orçamento é obrigatório para publicar uma defesa")
    if type_code == "defesa-orcamentaria" and defense_request is None:
        await _ensure_defense_open(year, scope_id_orcamento, scope_cloud, db)
    allow: set[int] = set()
    if defense_request is not None:
        departments, children = await _organization(db)
        allow = _subtree_ids(defense_request.department_id, children)
        prepared = [row for row in prepared if row["department_id"] in allow]
    else:
        departments, _ = await _organization(db)
        individual = await _individual_line_map(db, year, scope_id_orcamento, scope_cloud, departments) if type_code == "defesa-orcamentaria" else {}
        allow = set(individual.keys())
    prepared = await _scope_rows(prepared, scope_id_orcamento, scope_cloud, db, allow or None)
    if not prepared: raise HTTPException(422, "A carga não possui departamentos no escopo selecionado")
    snapshot = await _publish_version(fin_type, year, prepared, "IMPORT_EXCEL", db, title=title, description=description, event_type=event_type, effective_date=effective_date, filename=file.filename, scope_id_orcamento=scope_id_orcamento, scope_cloud=scope_cloud, defense_request_id=defense_request.id if defense_request else None)
    if defense_request is not None:
        defense_request.defense_version_id = snapshot.id
        if defense_request.status == "SOLICITADA":
            defense_request.status = "RASCUNHO"
        await db.commit()
    return {"version": _out(snapshot), "rows": len(prepared), "warnings": []}


async def _timeline(fin_type: FinancialType, year: int, db: AsyncSession) -> list[dict[str, Any]]:
    versions = (await db.execute(select(PlanningVersion).where(PlanningVersion.financial_type_id == fin_type.id, PlanningVersion.year == year).order_by(PlanningVersion.effective_date, PlanningVersion.created_at, PlanningVersion.version))).scalars().all()
    result: list[dict[str, Any]] = []
    previous_total: float | None = None
    cumulative_total = 0.0
    for version in versions:
        total = float((await db.execute(select(func.coalesce(func.sum(PlanningEntry.value), 0)).where(PlanningEntry.planning_version_id == version.id))).scalar_one())
        if fin_type.code == "orcamento":
            # Budget imports are additive ledger entries, not replacement
            # snapshots. Show the accumulated balance and the amount of this
            # specific load as the delta (negative loads are cuts).
            delta = total
            cumulative_total += total
            annual_total = cumulative_total
            delta_pct = None if previous_total in (None, 0) else delta / previous_total * 100
            previous_total = annual_total
        else:
            delta = None if previous_total is None else total - previous_total
            annual_total = total
            delta_pct = None if previous_total in (None, 0) else delta / previous_total * 100
            previous_total = total
        result.append({**_out(version), "title": version.title or f"Versão legada v{version.version}", "description": version.description or "Snapshot sem narrativa cadastrada.", "event_type": version.event_type or "OUTRO", "effective_date": version.effective_date or version.created_at.date(), "annual_total": annual_total, "delta": delta, "delta_pct": delta_pct})
    return result


def _timeline_event_matches_budget_scope(event: dict[str, Any], scoped_budget_ids: set[str] | None) -> bool:
    """Return whether a snapshot can affect the selected organizational scope.

    Older and organization-wide budget snapshots do not carry an explicit
    ``scope_id_orcamento``.  They still contain entries for every department,
    so excluding them from a department-filtered dashboard hides their changes
    and breaks the historical delta chain.
    """
    if scoped_budget_ids is None:
        return True
    event_budget_id = str(event.get("scope_id_orcamento") or "")
    return not event_budget_id or event_budget_id in scoped_budget_ids


@router.get("/dashboard/overview")
async def dashboard_overview(year: int, type_codes: str | None = None, department_id: int | None = None, db: AsyncSession = Depends(get_db)):
    codes = [code for code in (type_codes.split(",") if type_codes else ["orcamento", "forecast-manual", "base-line-contratual", "carta-credito", "defesa-orcamentaria"]) if code]
    types = (await db.execute(select(FinancialType).where(FinancialType.code.in_(codes)))).scalars().all()
    type_by_code = {item.code: item for item in types}
    departments, children = await _organization(db)
    scoped_department_ids: set[int] | None = None
    scope_path: list[str] | None = None
    if department_id is not None:
        if department_id not in departments:
            raise HTTPException(404, "Departamento não encontrado")
        scoped_department_ids = set()
        pending = [department_id]
        while pending:
            current_id = pending.pop()
            if current_id in scoped_department_ids:
                continue
            scoped_department_ids.add(current_id)
            pending.extend(child.id for child in children.get(current_id, []))
        scope_path = _levels(departments[department_id], departments)
    scoped_clouds = {
        str(_department_scope(departments[item_id], departments)[1] or "").lower()
        for item_id in (scoped_department_ids or set())
    }
    # A seleção de um nó do organograma representa o orçamento herdado por
    # esse nó. Não podemos unir os IDs próprios de todos os descendentes: isso
    # faria, por exemplo, uma defesa EXTERNOS aparecer ao selecionar Azure.
    selected_budget_id = (
        str(_department_scope(departments[department_id], departments)[0] or "")
        if department_id is not None else ""
    )
    scoped_budget_ids = {selected_budget_id} if selected_budget_id else {
        budget_id
        for item_id in (scoped_department_ids or set())
        if (budget_id := str(_department_scope(departments[item_id], departments)[0] or ""))
    }

    async def scoped_total(version_id: int) -> float:
        entries = (await db.execute(select(PlanningEntry).where(PlanningEntry.planning_version_id == version_id))).scalars().all()
        return sum(float(entry.value or 0) for entry in entries if scoped_department_ids is None or entry.department_id in scoped_department_ids)

    cards, series, events = [], {}, []
    for code in codes:
        fin_type = type_by_code.get(code)
        if not fin_type: continue
        # ``year`` é sempre o exercício financiado. A defesa de 2026 pertence
        # a 2026, embora suas decisões tenham effective_date em 2025.
        timeline_year = year
        timeline = await _timeline(fin_type, timeline_year, db)
        if not timeline: continue
        # Há uma versão vigente por ID de orçamento. O cartão deve consolidar
        # todas as versões vigentes que pertencem ao escopo, e não escolher
        # arbitrariamente a última versão de outro ID.
        scoped_timeline = [
            row for row in timeline
            if _timeline_event_matches_budget_scope(
                row,
                None if scoped_department_ids is None else scoped_budget_ids,
            )
        ]
        if not scoped_timeline:
            continue
        current_versions = [row for row in scoped_timeline if row["is_current"]]
        if not current_versions:
            current_versions = [scoped_timeline[-1]]
        current = max(current_versions, key=lambda row: (str(row.get("effective_date") or ""), row["id"]))
        entries = []
        for version in current_versions:
            entries.extend((await db.execute(select(PlanningEntry).where(PlanningEntry.planning_version_id == version["id"]))).scalars().all())
        if scoped_department_ids is not None:
            entries = [entry for entry in entries if entry.department_id in scoped_department_ids]
        monthly = [0.0] * 12
        for entry in entries: monthly[entry.month - 1] += float(entry.value or 0)
        cards.append({"code": code, "name": fin_type.name, "current": current, "annual_total": sum(monthly), "monthly": monthly})
        series[code] = monthly
        previous_total: float | None = None
        cumulative_event_total = 0.0
        for event in scoped_timeline:
            event_cloud = str(event.get("scope_cloud") or "").lower()
            event_budget_id = str(event.get("scope_id_orcamento") or "")
            # Um evento restrito a uma cloud não deve construir uma narrativa
            # artificial em outra. Ex.: defesa/inclusão Azure não aparece em GCP.
            if event_cloud and scoped_department_ids is not None and event_cloud not in scoped_clouds:
                continue
            if event_budget_id and scoped_department_ids is not None and event_budget_id not in scoped_budget_ids:
                continue
            event_amount = await scoped_total(event["id"])
            if code == "orcamento":
                cumulative_event_total += event_amount
                event_total = cumulative_event_total
                event_delta = event_amount
            else:
                event_total = event_amount
                event_delta = None if previous_total is None else event_total - previous_total
            # Eventos de inclusão, aporte ou corte só são relevantes para um
            # nó quando mudam seu orçamento. Isso esconde a inclusão de Aura
            # Pessoas ao navegar por uma diretoria sem esse projeto.
            if event_delta is not None and abs(event_delta) < 0.005 and event.get("event_type") in {"CORTE", "INCLUSAO_PROJETO", "AJUSTE", "REALOCACAO"}:
                continue
            scoped_event = {
                **event,
                "annual_total": event_total,
                "delta": event_delta,
                "delta_pct": None if previous_total in (None, 0) else event_delta / previous_total * 100,
                "type_code": code,
                "type_name": fin_type.name,
                "planning_year": timeline_year,
                "route": {"orcamento": "/budgets", "forecast-manual": "/manual-forecast", "base-line-contratual": "/contractual-baseline", "carta-credito": "/credit-letter", "defesa-orcamentaria": "/budget-defense"}.get(code, "/budgets"),
            }
            events.append(scoped_event)
            previous_total = event_total
    events.sort(key=lambda event: (str(event["effective_date"]), str(event["created_at"])), reverse=True)
    return {"year": year, "cards": cards, "series": series, "events": events, "scope": {"department_id": department_id, "path": scope_path, "included_departments": len(scoped_department_ids or departments)}}


@router.get("/budget-accuracy")
async def budget_accuracy(
    year: int,
    month: int | None = None,
    period: str = "monthly",
    department_id: int | None = None,
    scope_id_orcamento: str | None = None,
    scope_cloud: str | None = None,
    db: AsyncSession = Depends(get_db),
):
    if period not in {"monthly", "annual"}:
        raise HTTPException(422, "period deve ser monthly ou annual")
    if period == "monthly" and (month is None or month < 1 or month > 12):
        raise HTTPException(422, "O mês deve estar entre 1 e 12")
    budget_type = (await db.execute(select(FinancialType).where(FinancialType.code == "orcamento"))).scalar_one_or_none()
    if not budget_type:
        raise HTTPException(404, "Tipificação de orçamento não encontrada")
    departments, children = await _organization(db)
    if department_id is not None and department_id not in departments:
        raise HTTPException(404, "Departamento não encontrado")
    included_ids = _subtree_ids(department_id, children) if department_id is not None else None
    if department_id is not None:
        inherited_id, inherited_cloud = _department_scope(departments[department_id], departments)
        scope_id_orcamento = scope_id_orcamento or inherited_id
        scope_cloud = scope_cloud or inherited_cloud

    base_query = select(PlanningVersion).where(
        PlanningVersion.financial_type_id == budget_type.id,
        PlanningVersion.year == year,
        PlanningVersion.is_current == True,
    )
    query = base_query
    if scope_id_orcamento:
        query = query.where(PlanningVersion.scope_id_orcamento == scope_id_orcamento)
    if scope_cloud:
        query = query.where(PlanningVersion.scope_cloud == scope_cloud)
    versions = list((await db.execute(query.order_by(PlanningVersion.effective_date.desc(), PlanningVersion.version.desc()))).scalars().all())
    if not versions and (scope_id_orcamento or scope_cloud):
        fallback = base_query.where(PlanningVersion.scope_cloud.is_(None))
        if scope_id_orcamento:
            fallback = fallback.where(or_(
                PlanningVersion.scope_id_orcamento == scope_id_orcamento,
                PlanningVersion.scope_id_orcamento.is_(None),
            ))
        versions = list((await db.execute(fallback.order_by(PlanningVersion.effective_date.desc(), PlanningVersion.version.desc()))).scalars().all())

    selected_months = list(range(1, 13)) if period == "annual" else [month]
    budget_by_department: dict[int, float] = defaultdict(float)
    if versions:
        entries = (await db.execute(select(PlanningEntry).where(
            PlanningEntry.planning_version_id.in_([version.id for version in versions]),
            PlanningEntry.entity_type == "department",
            PlanningEntry.month.in_(selected_months),
        ))).scalars().all()
        for entry in entries:
            if entry.department_id is not None and (included_ids is None or entry.department_id in included_ids):
                budget_by_department[entry.department_id] += float(entry.value or 0)

    actual_by_department: dict[int, float] = defaultdict(float)
    forecast_by_department: dict[int, float] = defaultdict(float)
    actual_total = 0.0
    forecast_total = 0.0
    current_month_key = datetime.now().strftime("%Y%m")
    path_to_department: dict[tuple[str, ...], int] = {}
    for department in departments.values():
        if not department.active:
            continue
        path = _levels(department, departments)
        path_to_department[tuple(path)] = department.id
        if len(path) >= 4:
            director_path = [path[index] for index in (0, 1, 3, 2, *range(4, len(path)))]
            path_to_department.setdefault(tuple(director_path), department.id)
    inherited_cloud = _department_scope(departments[department_id], departments)[1] if department_id is not None else None
    selected_path = _levels(departments[department_id], departments) if department_id is not None else None
    actual_rows = await _budget_accuracy_actuals(year, selected_months, inherited_cloud, selected_path)
    consolidate_selected_project = bool(
        selected_path and len(selected_path) >= 5 and selected_path[1] in {"Showback", "Chargeback"}
    )
    for row in actual_rows:
        period_key = str(row.get("period") or "")
        actual_value = float(row.get("actual_value") or 0)
        forecast_value = float(row.get("forecast_value") or 0)
        actual_total += actual_value
        if not row.get("has_actual") or period_key == current_month_key:
            forecast_total += forecast_value
        path = tuple(str(row.get(f"organograma_nivel_{level}") or "").strip() for level in range(1, 8))
        path = tuple(item for item in path if item)
        current_id = path_to_department.get(path)
        if consolidate_selected_project and department_id is not None:
            selected_tail = tuple(selected_path[2:])
            natural_tail = tuple(path[2:]) if len(path) >= 5 else ()
            swapped_tail = (path[3], path[2], *path[4:]) if len(path) >= 5 else ()
            if natural_tail == selected_tail or swapped_tail == selected_tail:
                current_id = department_id
        if current_id is not None and (included_ids is None or current_id in included_ids):
            actual_by_department[current_id] += actual_value
            if not row.get("has_actual") or period_key == current_month_key:
                forecast_by_department[current_id] += forecast_value

    line_ids = set(budget_by_department) | set(actual_by_department) | set(forecast_by_department)
    lines = []
    for current_id in line_ids:
        department = departments.get(current_id)
        if not department:
            continue
        budget_value = budget_by_department.get(current_id, 0.0)
        actual_value = actual_by_department.get(current_id, 0.0)
        forecast_value = forecast_by_department.get(current_id, 0.0)
        comparison_value = actual_value + forecast_value if period == "annual" else actual_value
        lines.append({
            "department_id": current_id,
            "department_levels": _levels(department, departments),
            **_budget_accuracy_metrics(budget_value, comparison_value),
            "actual_value": actual_value,
            "forecast_value": forecast_value,
            "comparison_value": comparison_value,
        })
    if period == "annual":
        for line in lines:
            line["variance"] = -line["variance"]
            if line["variance_pct"] is not None:
                line["variance_pct"] = -line["variance_pct"]
        lines.sort(key=lambda item: item["variance"])
    else:
        lines.sort(key=lambda item: item["variance"], reverse=True)
    totals = _budget_accuracy_metrics(
        sum(item["budget_value"] for item in lines),
        actual_total + forecast_total if period == "annual" else actual_total,
    )
    totals["actual_value"] = actual_total
    totals["forecast_value"] = forecast_total
    totals["comparison_value"] = actual_total + forecast_total if period == "annual" else actual_total
    if period == "annual":
        totals["variance"] = -totals["variance"]
        if totals["variance_pct"] is not None:
            totals["variance_pct"] = -totals["variance_pct"]
    selected_version = max(versions, key=lambda item: (item.effective_date or date.min, item.version)) if versions else None
    return {
        "year": year,
        "period": period,
        "month": month,
        "target_month": f"{year}{month:02d}" if month else None,
        "version": _out(selected_version) if selected_version else None,
        "scope": {"department_id": department_id, "path": _levels(departments[department_id], departments) if department_id else None},
        "totals": totals,
        "lines": lines,
    }


@router.get("/budget-accuracy/detail")
async def budget_accuracy_detail(
    year: int,
    department_id: int,
    db: AsyncSession = Depends(get_db),
):
    departments, _ = await _organization(db)
    department = departments.get(department_id)
    if not department or not department.active:
        raise HTTPException(404, "Departamento não encontrado")
    path = _levels(department, departments)
    inherited_cloud = _department_scope(department, departments)[1]
    periods = [f"{year}{month:02d}" for month in range(1, 13)]
    escaped_path = [item.strip().replace("'", "''") for item in path]
    candidates = [list(range(1, len(escaped_path) + 1))]
    if len(escaped_path) >= 3:
        candidates.append([1, 2, 4, 3, *range(5, len(escaped_path) + 2)])
    if len(escaped_path) >= 5 and escaped_path[1] in {"Showback", "Chargeback"}:
        candidates.extend([
            [1, 3, 4, *range(5, len(escaped_path) + 2)],
            [1, 4, 3, *range(5, len(escaped_path) + 2)],
        ])
    scope_checks = []
    for order in candidates:
        scope_checks.append(" AND ".join(
            f'TRIM(CAST("organograma_nivel_{physical}" AS VARCHAR)) = \'{escaped_path[index]}\''
            for index, physical in enumerate(order[:len(escaped_path)])
        ))
    scope_sql = " AND (" + " OR ".join(f"({item})" for item in scope_checks) + ")"
    cloud_sql = ""
    if inherited_cloud:
        cloud = inherited_cloud.lower()
        cloud_sql = f" AND lower(CAST(\"cloud_origem\" AS VARCHAR)) LIKE '{cloud}%'"
    period_sql = ", ".join(f"'{item}'" for item in periods)
    try:
        rows = await trino_execute(f"""
                 SELECT CAST(period AS VARCHAR) AS period, service_subcategory,
                     SUM(actual_value) AS actual_value, SUM(forecast_value) AS forecast_value, MAX(has_actual) AS has_actual
            FROM (
              SELECT CAST("_billing_period" AS VARCHAR) AS period,
                     COALESCE(NULLIF(TRIM(CAST("servicesubcategory" AS VARCHAR)), ''), 'Outros') AS service_subcategory,
                     SUM(COALESCE(TRY_CAST("realizado" AS DOUBLE), 0)) AS actual_value,
                     CAST(0 AS DOUBLE) AS forecast_value, 1 AS has_actual
              FROM iceberg.{GOLD}.dest_dm_cost_visao_completa
              WHERE CAST("_billing_period" AS VARCHAR) IN ({period_sql})
                AND ("tipo" IS NULL OR TRIM(CAST("tipo" AS VARCHAR)) IN ('', 'realizado')){cloud_sql}{scope_sql}
              GROUP BY 1, 2
              UNION ALL
              SELECT CAST("period" AS VARCHAR) AS period,
                     COALESCE(NULLIF(TRIM(CAST("servicesubcategory" AS VARCHAR)), ''), 'Outros') AS service_subcategory,
                     CAST(0 AS DOUBLE) AS actual_value,
                     SUM(COALESCE(TRY_CAST("realizado" AS DOUBLE), 0)) AS forecast_value, 0 AS has_actual
              FROM iceberg.{GOLD}.dest_dm_cost_visao_completa
              WHERE CAST("period" AS VARCHAR) IN ({period_sql})
                AND lower(TRIM(CAST("tipo" AS VARCHAR))) = 'forecast'{cloud_sql}{scope_sql}
              GROUP BY 1, 2
            ) facts
            GROUP BY period, service_subcategory
            ORDER BY period, actual_value DESC, forecast_value DESC
        """, timeout_s=180)
    except Exception as exc:
        raise HTTPException(503, f"Falha ao consultar detalhe do projeto: {exc}")

    budget_rows = (await db.execute(select(PlanningEntry).where(
        PlanningEntry.department_id == department_id,
        PlanningEntry.entity_type == "department",
        PlanningEntry.month.in_(range(1, 13)),
        PlanningEntry.planning_version_id.in_(select(PlanningVersion.id).where(
            PlanningVersion.financial_type_id == select(FinancialType.id).where(FinancialType.code == "orcamento").scalar_subquery(),
            PlanningVersion.year == year,
            PlanningVersion.is_current == True,
        )),
    ))).scalars().all()
    budget_by_month = defaultdict(float)
    for entry in budget_rows:
        budget_by_month[entry.month] += float(entry.value or 0)
    monthly = {period: {"period": period, "budget_value": budget_by_month[index + 1], "actual_value": 0.0, "forecast_value": 0.0} for index, period in enumerate(periods)}
    subcategories: dict[str, float] = defaultdict(float)
    subcategory_monthly: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    current_month_key = datetime.now().strftime("%Y%m")
    for row in rows:
        item = monthly.setdefault(str(row["period"]), {"period": str(row["period"]), "budget_value": 0.0, "actual_value": 0.0, "forecast_value": 0.0})
        actual_value = float(row.get("actual_value") or 0)
        forecast_value = float(row.get("forecast_value") or 0) if (not row.get("has_actual") or str(row["period"]) == current_month_key) else 0.0
        item["actual_value"] += actual_value
        item["forecast_value"] += forecast_value
        subcategory = str(row.get("service_subcategory") or "Outros")
        consumption = actual_value + forecast_value
        subcategories[subcategory] += consumption
        subcategory_monthly[subcategory][str(row["period"])] += consumption
    ranked_subcategories = sorted(subcategories.items(), key=lambda item: item[1], reverse=True)[:5]
    top_names = {name for name, _ in ranked_subcategories}
    subcategory_series = [
        {
            "name": name,
            "value": value,
            "monthly_values": [subcategory_monthly[name].get(period, 0.0) for period in periods],
        }
        for name, value in ranked_subcategories
    ]
    others_monthly = [
        sum(values.get(period, 0.0) for name, values in subcategory_monthly.items() if name not in top_names)
        for period in periods
    ]
    return {
        "year": year,
        "department_id": department_id,
        "department_levels": path,
        "monthly": list(monthly.values()),
        "subcategories": subcategory_series,
        "others_value": sum(subcategories.values()) - sum(value for _, value in ranked_subcategories),
        "others_monthly": others_monthly,
    }


@router.post("/budget-accuracy/acks", status_code=201)
async def create_budget_accuracy_ack(body: BudgetAccuracyAckIn, db: AsyncSession = Depends(get_db)):
    if not re.fullmatch(r"\d{6}", body.target_month):
        raise HTTPException(422, "target_month deve estar no formato YYYYMM")
    path = [part.strip() for part in body.department.split(">") if part.strip()]
    if not path:
        raise HTTPException(422, "Informe o departamento")
    row = BudgetAccuracyAck(
        target_month=body.target_month,
        planning_version_id=body.planning_version_id,
        department_id=body.department_id,
        department_path=path,
        department_name=path[-1],
        budget_value=body.budget_value,
        actual_value=body.actual_value,
        variance=body.actual_value - body.budget_value,
        category=body.category.strip() or "Outro",
        reason=body.reason.strip(),
        severity=body.severity.strip() or "Informação",
        ordem=body.ordem,
        registered_at=body.registered_at,
    )
    db.add(row)
    await db.commit()
    await db.refresh(row)
    return _out(row)


@router.get("/budget-accuracy/acks")
async def list_budget_accuracy_acks(
    target_month: str | None = None,
    department_id: int | None = None,
    db: AsyncSession = Depends(get_db),
):
    query = select(BudgetAccuracyAck).where(BudgetAccuracyAck.active == True).order_by(BudgetAccuracyAck.created_at.desc())
    if target_month:
        query = query.where(BudgetAccuracyAck.target_month == target_month)
    if department_id is not None:
        query = query.where(BudgetAccuracyAck.department_id == department_id)
    rows = (await db.execute(query.limit(500))).scalars().all()
    return {"items": [_out(row) for row in rows]}


def _group_entries(entries: list[PlanningEntry]) -> dict[int, dict[str, Any]]:
    grouped: dict[int, dict[str, Any]] = {}
    for entry in entries:
        if entry.department_id is None:
            continue
        grouped.setdefault(entry.department_id, {"entity_type": "department", "department_id": entry.department_id, "department_levels": entry.department_levels, "values": [0.0] * 12})["values"][entry.month - 1] = float(entry.value or 0)
    return grouped


async def _defense_totals(db: AsyncSession, requests: list[BudgetDefenseRequest]) -> dict[int, float]:
    version_ids = [row.defense_version_id for row in requests if row.defense_version_id]
    if not version_ids:
        return {}
    rows = (await db.execute(
        select(PlanningEntry.planning_version_id, func.coalesce(func.sum(PlanningEntry.value), 0))
        .where(PlanningEntry.planning_version_id.in_(version_ids))
        .group_by(PlanningEntry.planning_version_id)
    )).all()
    return {int(version_id): float(total or 0) for version_id, total in rows}


@router.get("/budget-defense/requests")
async def list_defense_requests(
    year: int,
    status: str | None = None,
    department_id: int | None = None,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require("planning", "read")),
):
    """Solicitações de defesa do exercício, com status e linha solicitada."""
    query = select(BudgetDefenseRequest).where(BudgetDefenseRequest.year == year)
    if status:
        statuses = tuple(item.strip().upper() for item in status.split(",") if item.strip())
        if statuses:
            query = query.where(BudgetDefenseRequest.status.in_(statuses))
    if department_id:
        query = query.where(BudgetDefenseRequest.department_id == department_id)
    requests = list((await db.execute(query.order_by(BudgetDefenseRequest.id.desc()))).scalars().all())
    departments, _ = await _organization(db)
    totals = await _defense_totals(db, requests)
    return [
        {**_defense_request_out(row, departments), "annual_total": totals.get(row.defense_version_id, 0.0)}
        for row in requests
    ]


@router.get("/budget-defense/departments")
async def defense_department_options(
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require("planning", "read")),
):
    """Organograma (com ID Orçamento herdado) para ancorar projeto novo ou escolher um existente."""
    departments, _ = await _organization(db)
    items = []
    for department in sorted((item for item in departments.values() if item.active), key=lambda item: _levels(item, departments)):
        budget_id, cloud = _department_scope(department, departments)
        items.append({
            "id": department.id,
            "parent_id": department.parent_id,
            "name": department.name,
            "path": _levels(department, departments),
            "centro_custo": department.centro_custo,
            "id_orcamento": department.id_orcamento,
            "budget_id": budget_id,
            "cloud": cloud,
        })
    return items


@router.post("/budget-defense/projects", status_code=201)
async def create_defense_project(
    body: DefenseProjectIn,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require("planning", "create", DEFENSE_SCOPE_REQUEST)),
):
    """Cria o projeto no organograma e abre a solicitação de defesa dele."""
    departments, _ = await _organization(db)
    parent = departments.get(body.parent_id)
    if not parent or not parent.active:
        raise HTTPException(422, "Selecione um nó ativo do organograma para receber o projeto")
    if len(_levels(parent, departments)) >= 7:
        raise HTTPException(422, "A hierarquia suporta no máximo 7 níveis")
    name = (body.name or "").strip()
    if not name:
        raise HTTPException(422, "Informe o nome do projeto")
    new_id = (body.id_orcamento or "").strip()
    if not new_id:
        raise HTTPException(422, "Informe o novo ID Orçamento do projeto")
    duplicate = next((item for item in departments.values() if item.active and item.parent_id == parent.id and (item.name or "").strip().lower() == name.lower()), None)
    if duplicate:
        raise HTTPException(422, "Já existe um projeto com este nome neste nó do organograma")
    department = Department(parent_id=parent.id, name=name, id_orcamento=new_id, centro_custo=body.centro_custo, ambiente=body.ambiente, cloud_filter=body.cloud_filter, active=True)
    db.add(department)
    await db.flush()
    departments[department.id] = department
    cloud = _department_scope(parent, departments)[1]
    request = BudgetDefenseRequest(
        year=body.year,
        department_id=department.id,
        parent_department_id=parent.id,
        is_new_project=True,
        requested_id_orcamento=new_id,
        requested_cloud=cloud,
        title=body.title or f"Inclusão de projeto — {name}",
        description=body.description or "Projeto criado dentro da Defesa Orçamentária.",
        status="RASCUNHO",
        requested_by_id=user.id,
    )
    db.add(request)
    await db.commit()
    await db.refresh(department)
    await db.refresh(request)
    return {"department": _out(department), "request": _defense_request_out(request, departments)}


@router.post("/budget-defense/requests", status_code=201)
async def create_defense_request(
    body: DefenseRequestIn,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require("planning", "create", DEFENSE_SCOPE_REQUEST)),
):
    """Abre a defesa de um projeto existente — sem alterar o organograma."""
    department = await db.get(Department, body.department_id)
    if not department or not department.active:
        raise HTTPException(422, "Projeto inválido ou inativo no organograma")
    requested_id = (body.requested_id_orcamento or "").strip()
    if not requested_id:
        raise HTTPException(422, "Informe o ID Orçamento solicitado")
    departments, _ = await _organization(db)
    org_id, org_cloud = _department_scope(department, departments)
    existing = (await db.execute(
        select(BudgetDefenseRequest).where(
            BudgetDefenseRequest.year == body.year,
            BudgetDefenseRequest.department_id == department.id,
            BudgetDefenseRequest.status.in_(DEFENSE_OPEN_STATUSES),
        )
    )).scalars().all()
    if existing:
        raise HTTPException(422, "Este projeto já possui uma solicitação de defesa em andamento")
    request = BudgetDefenseRequest(
        year=body.year,
        department_id=department.id,
        parent_department_id=department.parent_id,
        is_new_project=False,
        requested_id_orcamento=requested_id,
        requested_cloud=(body.requested_cloud or org_cloud or None),
        title=body.title or f"Defesa de {department.name}",
        description=body.description,
        status="RASCUNHO",
        requested_by_id=user.id,
    )
    db.add(request)
    await db.commit()
    await db.refresh(request)
    return _defense_request_out(request, departments)


@router.put("/budget-defense/requests/{request_id}")
async def update_defense_request(
    request_id: int,
    body: DefenseRequestUpdateIn,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require("planning", "update", DEFENSE_SCOPE_REQUEST)),
):
    request = await _load_defense_request(db, request_id)
    if request.status not in ("RASCUNHO", "SOLICITADA"):
        raise HTTPException(422, "Solicitação já decidida; não é possível alterar a linha pedida")
    data = body.model_dump(exclude_unset=True)
    if "requested_id_orcamento" in data:
        value = (data["requested_id_orcamento"] or "").strip()
        if not value:
            raise HTTPException(422, "Informe o ID Orçamento solicitado")
        request.requested_id_orcamento = value
    if "requested_cloud" in data:
        request.requested_cloud = (data["requested_cloud"] or None)
    if "title" in data:
        request.title = data["title"]
    if "description" in data:
        request.description = data["description"]
    await db.commit()
    await db.refresh(request)
    departments, _ = await _organization(db)
    return _defense_request_out(request, departments)


@router.post("/budget-defense/requests/{request_id}/submit")
async def submit_defense_request(
    request_id: int,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require("planning", "update", DEFENSE_SCOPE_REQUEST)),
):
    request = await _load_defense_request(db, request_id)
    if request.status not in ("RASCUNHO", "SOLICITADA"):
        raise HTTPException(422, "Solicitação já decidida")
    if not request.defense_version_id:
        raise HTTPException(422, "Salve os valores da defesa antes de enviar para aprovação")
    request.status = "SOLICITADA"
    request.requested_by_id = request.requested_by_id or user.id
    await db.commit()
    await db.refresh(request)
    departments, _ = await _organization(db)
    totals = await _defense_totals(db, [request])
    return {**_defense_request_out(request, departments), "annual_total": totals.get(request.defense_version_id, 0.0)}


@router.post("/budget-defense/requests/{request_id}/approve")
async def approve_defense_request(
    request_id: int,
    body: DefenseApproveIn,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require("planning", "update", DEFENSE_SCOPE_PUBLISH)),
):
    """Aprova a defesa escolhendo o ID Orçamento final (mantém ou troca o pedido)."""
    request = await _load_defense_request(db, request_id)
    if request.status in ("PUBLICADA", "REJEITADA"):
        raise HTTPException(422, "Solicitação já encerrada")
    if request.status == "RASCUNHO":
        raise HTTPException(422, "A defesa ainda é rascunho; solicite o envio antes de aprovar")
    final_id = (body.final_id_orcamento or "").strip()
    if not final_id:
        raise HTTPException(422, "Informe o ID Orçamento final")
    note = (body.decision_note or "").strip()
    if not note:
        raise HTTPException(422, "Descreva o motivo e as observações da aprovação")
    request.decided_by_id = user.id
    request.decided_at = datetime.now(timezone.utc)
    request.final_id_orcamento = final_id
    request.final_cloud = body.final_cloud if body.final_cloud is not None else request.requested_cloud
    request.decision_note = note
    request.status = "APROVADA"
    await db.commit()
    await db.refresh(request)
    departments, _ = await _organization(db)
    totals = await _defense_totals(db, [request])
    return {**_defense_request_out(request, departments), "annual_total": totals.get(request.defense_version_id, 0.0)}


@router.post("/budget-defense/requests/{request_id}/reject")
async def reject_defense_request(
    request_id: int,
    body: DefenseRejectIn,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require("planning", "update", DEFENSE_SCOPE_PUBLISH)),
):
    request = await _load_defense_request(db, request_id)
    if request.status == "PUBLICADA":
        raise HTTPException(422, "Solicitação já publicada no orçamento")
    note = (body.decision_note or "").strip()
    if not note:
        raise HTTPException(422, "Descreva o motivo da rejeição")
    request.decided_by_id = user.id
    request.decided_at = datetime.now(timezone.utc)
    request.decision_note = note
    request.status = "REJEITADA"
    await db.commit()
    await db.refresh(request)
    departments, _ = await _organization(db)
    return _defense_request_out(request, departments)


@router.post("/budget-defense/requests/{request_id}/publish")
async def publish_defense_request(
    request_id: int,
    body: DefensePublishIn,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require("planning", "update", DEFENSE_SCOPE_PUBLISH)),
):
    """Publica a defesa no Orçamento e materializa o ID Orçamento no organograma.

    É aqui que a "linha individual" deixa de existir: o projeto passa a
    pertencer a um único ID Orçamento, o final escolhido na aprovação.
    """
    request = await _load_defense_request(db, request_id)
    if request.status == "PUBLICADA":
        raise HTTPException(422, "Solicitação já publicada no orçamento")
    if request.status == "REJEITADA":
        raise HTTPException(422, "Solicitação rejeitada não pode ser publicada")
    if request.status == "RASCUNHO":
        raise HTTPException(422, "Envie a defesa para aprovação antes de publicar no orçamento")
    if request.status == "SOLICITADA":
        raise HTTPException(422, "Aprove a defesa (definindo o ID Orçamento final) antes de publicar no orçamento")
    defense_type = (await db.execute(select(FinancialType).where(FinancialType.code == "defesa-orcamentaria"))).scalar_one_or_none()
    budget_type = (await db.execute(select(FinancialType).where(FinancialType.code == "orcamento"))).scalar_one_or_none()
    if not defense_type or not budget_type:
        raise HTTPException(404, "Tipificações financeiras não encontradas")
    defense_version = await db.get(PlanningVersion, request.defense_version_id) if request.defense_version_id else None
    if not defense_version:
        defense_version = (await db.execute(
            select(PlanningVersion).where(
                PlanningVersion.defense_request_id == request.id,
                PlanningVersion.financial_type_id == defense_type.id,
            ).order_by(PlanningVersion.version.desc())
        )).scalars().first()
    if not defense_version:
        raise HTTPException(422, "Salve a defesa antes de publicar no orçamento")
    final_id = (body.final_id_orcamento or request.final_id_orcamento or request.requested_id_orcamento or "").strip()
    if not final_id:
        raise HTTPException(422, "Informe o ID Orçamento final")
    final_cloud = body.final_cloud if body.final_cloud is not None else (request.final_cloud or request.requested_cloud)
    departments, children = await _organization(db)
    department = departments.get(request.department_id)
    if not department:
        raise HTTPException(404, "Projeto da solicitação não existe mais no organograma")
    subtree = _subtree_ids(request.department_id, children)
    defense_entries = (await db.execute(
        select(PlanningEntry).where(
            PlanningEntry.planning_version_id == defense_version.id,
            PlanningEntry.department_id.in_(subtree),
        )
    )).scalars().all()
    if not defense_entries:
        raise HTTPException(422, "A defesa não possui valores salvos para este projeto")
    # Consolida o orçamento vigente e substitui somente o subtree do projeto.
    budget_versions = (await db.execute(
        select(PlanningVersion).where(
            PlanningVersion.financial_type_id == budget_type.id,
            PlanningVersion.year == request.year,
            PlanningVersion.is_current == True,
        )
    )).scalars().all()
    budget_entries: list[PlanningEntry] = []
    for version in budget_versions:
        budget_entries.extend((await db.execute(select(PlanningEntry).where(PlanningEntry.planning_version_id == version.id))).scalars().all())
    merged = _group_entries(budget_entries)
    for department_id in subtree:
        merged.pop(department_id, None)
    merged.update(_group_entries(defense_entries))
    snapshot = await _publish_version(
        budget_type, request.year, list(merged.values()), "DEFESA_APROVADA", db,
        title=body.title, description=body.description, event_type="APROVACAO",
        effective_date=body.effective_date, scope_id_orcamento=final_id, scope_cloud=final_cloud,
    )
    now = datetime.now(timezone.utc)
    # Materializa a linha: o projeto passa a pertencer ao ID Orçamento final.
    department.id_orcamento = final_id
    if final_cloud and not department.cloud_filter:
        department.cloud_filter = final_cloud
    request.status = "PUBLICADA"
    request.final_id_orcamento = final_id
    request.final_cloud = final_cloud
    request.published_budget_version_id = snapshot.id
    request.published_at = now
    request.decided_by_id = request.decided_by_id or user.id
    request.decided_at = request.decided_at or now
    if body.decision_note:
        request.decision_note = body.decision_note.strip()
    defense_version.published_to_budget_at = now
    defense_version.published_budget_version_id = snapshot.id
    await db.commit()
    await db.refresh(request)
    await db.refresh(department)
    return {
        "request": _defense_request_out(request, departments),
        "budget_version": _out(snapshot),
        "department": _out(department),
        "rows": len(defense_entries),
    }


@router.get("/{type_code}/{year}")
async def planning_grid(type_code: str, year: int, version_id: int | None = None, scope_id_orcamento: str | None = None, scope_cloud: str | None = None, defense_request_id: int | None = None, db: AsyncSession = Depends(get_db)):
    fin_type = (await db.execute(select(FinancialType).where(FinancialType.code == type_code))).scalar_one_or_none()
    if not fin_type: raise HTTPException(404, "Tipificação financeira não encontrada")
    defense_request: BudgetDefenseRequest | None = None
    if defense_request_id and type_code == "defesa-orcamentaria":
        candidate = await db.get(BudgetDefenseRequest, defense_request_id)
        if candidate and candidate.year == year:
            defense_request = candidate
            # A grade de uma solicitação sempre usa a linha pedida por ela.
            scope_id_orcamento = candidate.requested_id_orcamento
            scope_cloud = candidate.requested_cloud
    base_query = select(PlanningVersion).where(PlanningVersion.financial_type_id == fin_type.id, PlanningVersion.year == year)
    query = base_query
    if version_id: query = query.where(PlanningVersion.id == version_id)
    elif defense_request is not None:
        query = query.where(PlanningVersion.is_current == True, PlanningVersion.defense_request_id == defense_request.id)
    else:
        query = query.where(PlanningVersion.is_current == True)
        if scope_id_orcamento: query = query.where(PlanningVersion.scope_id_orcamento == scope_id_orcamento)
        if scope_cloud: query = query.where(PlanningVersion.scope_cloud == scope_cloud)
    versions = (await db.execute(query.order_by(PlanningVersion.effective_date.desc(), PlanningVersion.version.desc()))).scalars().all()
    # Snapshots anteriores ao recorte por Cloud guardam o ID Orçamento, mas
    # possuem scope_cloud nulo. Ao selecionar Azure/GCP, use esse snapshot como
    # fallback e restrinja suas linhas pela hierarquia logo abaixo.
    if not versions and not version_id and defense_request is None and (scope_id_orcamento or scope_cloud):
        fallback_query = base_query.where(
            PlanningVersion.is_current == True,
            PlanningVersion.scope_cloud.is_(None),
        )
        if scope_id_orcamento:
            fallback_query = fallback_query.where(or_(
                PlanningVersion.scope_id_orcamento == scope_id_orcamento,
                PlanningVersion.scope_id_orcamento.is_(None),
            ))
        versions = (await db.execute(fallback_query.order_by(PlanningVersion.effective_date.desc(), PlanningVersion.version.desc()))).scalars().all()
    departments, children = await _organization(db)
    individual = await _individual_line_map(db, year, scope_id_orcamento, scope_cloud, departments) if type_code == "defesa-orcamentaria" else {}
    allow: set[int] = set(individual.keys())
    if defense_request is not None:
        # A solicitação pode acontecer fora da linha herdada; inclua o projeto.
        allow |= _subtree_ids(defense_request.department_id, children)
    payload: dict[str, Any] = {"type": _out(fin_type), "version": None, "rows": []}
    if type_code == "defesa-orcamentaria":
        payload["individual_lines"] = [_defense_request_out(item, departments) for item in individual.values()]
        payload["defense_request"] = _defense_request_out(defense_request, departments) if defense_request is not None else None
    if not versions:
        return payload
    # Orçamento/defesa podem ter uma versão vigente por ID de orçamento. Sem
    # filtro de ID, a grade deve compor todas elas; selecionar apenas a primeira
    # fazia os demais projetos parecerem zerados.
    version = versions[0]
    entries = (await db.execute(select(PlanningEntry).where(PlanningEntry.planning_version_id.in_([item.id for item in versions]), PlanningEntry.entity_type == "department").order_by(PlanningEntry.department_id, PlanningEntry.month))).scalars().all()
    grouped: dict[tuple[str, int, int | None], dict[str, Any]] = {}
    for entry in entries:
        key = (entry.entity_type, entry.department_id, None)
        # Multiple current budget loads are additive; compose them instead of
        # letting the last version overwrite the previous value in the grid.
        grouped.setdefault(key, {"entity_type": entry.entity_type, "department_id": entry.department_id, "department_levels": entry.department_levels, "values": [0.0] * 12})["values"][entry.month - 1] += float(entry.value or 0)
    scoped_rows = await _scope_rows(list(grouped.values()), scope_id_orcamento, scope_cloud, db, allow or None)
    payload["version"] = _out(version)
    payload["rows"] = scoped_rows
    return payload


@router.post("/{type_code}/{year}")
async def save_planning(type_code: str, year: int, body: PlanningSaveIn, db: AsyncSession = Depends(get_db), user: User = Depends(get_current_user)):
    fin_type = await _type_for_write(type_code, db)
    defense_request: BudgetDefenseRequest | None = None
    scope_id_orcamento = body.scope_id_orcamento
    scope_cloud = body.scope_cloud
    if body.defense_request_id:
        if type_code != "defesa-orcamentaria":
            raise HTTPException(422, "Solicitação de defesa só se aplica à Defesa Orçamentária")
        defense_request = await _load_defense_request(db, body.defense_request_id, year)
        if defense_request.status not in ("RASCUNHO", "SOLICITADA"):
            raise HTTPException(422, "Esta solicitação não aceita mais alterações de valores")
        scope_id_orcamento = defense_request.requested_id_orcamento
        scope_cloud = defense_request.requested_cloud
    prepared = await _validate_rows(body.rows, db)
    if type_code == "defesa-orcamentaria" and not scope_id_orcamento:
        raise HTTPException(422, "ID Orçamento é obrigatório para salvar uma defesa")
    if type_code == "defesa-orcamentaria" and defense_request is None:
        await _ensure_defense_open(year, scope_id_orcamento, scope_cloud, db)
    allow: set[int] = set()
    if defense_request is not None:
        departments, children = await _organization(db)
        # A defesa de uma solicitação grava somente o projeto solicitado; a
        # linha individual existe no organograma apenas como leitura.
        allow = _subtree_ids(defense_request.department_id, children)
        if not any(row["department_id"] in allow for row in prepared):
            raise HTTPException(422, "Inclua o projeto da solicitação na grade antes de salvar a defesa")
        prepared = [row for row in prepared if row["department_id"] in allow]
    else:
        departments, _ = await _organization(db)
        individual = await _individual_line_map(db, year, scope_id_orcamento, scope_cloud, departments) if type_code == "defesa-orcamentaria" else {}
        allow = set(individual.keys())
    prepared = await _scope_rows(prepared, scope_id_orcamento, scope_cloud, db, allow or None)
    if not prepared: raise HTTPException(422, "A grade não possui departamentos no escopo selecionado")
    if type_code == "orcamento":
        # The grid submits the current accumulated state. Since budget imports
        # are additive, an edit must publish only the difference from the
        # current ledger; otherwise the entire grid would be added again.
        current_version_ids = list((await db.execute(
            select(PlanningVersion.id).where(
                PlanningVersion.financial_type_id == fin_type.id,
                PlanningVersion.year == year,
                PlanningVersion.is_current == True,
            )
        )).scalars().all())
        current_values: dict[int, list[float]] = {}
        if current_version_ids:
            current_entries = (await db.execute(
                select(PlanningEntry).where(
                    PlanningEntry.planning_version_id.in_(current_version_ids),
                    PlanningEntry.entity_type == "department",
                )
            )).scalars().all()
            for entry in current_entries:
                values = current_values.setdefault(entry.department_id, [0.0] * 12)
                values[entry.month - 1] += float(entry.value or 0)
        for row in prepared:
            previous = current_values.get(row["department_id"], [0.0] * 12)
            row["values"] = [float(value) - previous[index] for index, value in enumerate(row["values"])]
    snapshot = await _publish_version(fin_type, year, prepared, "EDICAO_GRID", db, title=body.title, description=body.description, event_type=body.event_type, effective_date=body.effective_date, scope_id_orcamento=scope_id_orcamento, scope_cloud=scope_cloud, defense_request_id=defense_request.id if defense_request else None)
    if defense_request is not None:
        defense_request.defense_version_id = snapshot.id
        if not defense_request.requested_by_id:
            defense_request.requested_by_id = user.id
        if body.submit and defense_request.status == "RASCUNHO":
            defense_request.status = "SOLICITADA"
        elif not body.submit and defense_request.status == "SOLICITADA":
            # Editar valores de uma solicitação já enviada volta para rascunho.
            defense_request.status = "RASCUNHO"
        await db.commit()
        await db.refresh(defense_request)
    return {"version": _out(snapshot), "rows": len(prepared), "defense_request": _defense_request_out(defense_request, departments) if defense_request else None}


@router.get("/{type_code}/{year}/versions")
async def list_versions(type_code: str, year: int, db: AsyncSession = Depends(get_db)):
    fin_type = (await db.execute(select(FinancialType).where(FinancialType.code == type_code))).scalar_one_or_none()
    if not fin_type: raise HTTPException(404, "Tipificação financeira não encontrada")
    rows = (await db.execute(select(PlanningVersion).where(PlanningVersion.financial_type_id == fin_type.id, PlanningVersion.year == year).order_by(PlanningVersion.version.desc()))).scalars().all()
    return [_out(r) for r in rows]


@router.get("/{type_code}/{year}/timeline")
async def planning_timeline(type_code: str, year: int, db: AsyncSession = Depends(get_db)):
    fin_type = (await db.execute(select(FinancialType).where(FinancialType.code == type_code))).scalar_one_or_none()
    if not fin_type: raise HTTPException(404, "Tipificação financeira não encontrada")
    return await _timeline(fin_type, year, db)




@router.post("/defesa-orcamentaria/{year}/versions/{version_id}/publish-budget")
async def publish_defense_to_budget(year: int, version_id: int, body: BudgetPublishIn, db: AsyncSession = Depends(get_db)):
    defense_type = (await db.execute(select(FinancialType).where(FinancialType.code == "defesa-orcamentaria"))).scalar_one_or_none()
    budget_type = (await db.execute(select(FinancialType).where(FinancialType.code == "orcamento"))).scalar_one_or_none()
    defense = await db.get(PlanningVersion, version_id)
    if not defense_type or not budget_type or not defense or defense.financial_type_id != defense_type.id or defense.year != year:
        raise HTTPException(404, "Defesa orçamentária não encontrada")
    if defense.published_to_budget_at:
        raise HTTPException(422, "Esta defesa já foi carregada no orçamento; novas alterações devem ser revisões orçamentárias")
    await _ensure_defense_open(year, body.id_orcamento, body.cloud_scope, db)
    departments, _ = await _organization(db)
    scope_ids = {item.id for item in departments.values() if _department_scope(item, departments) == (body.id_orcamento, body.cloud_scope) or (body.cloud_scope is None and _department_scope(item, departments)[0] == body.id_orcamento)}
    if not scope_ids:
        raise HTTPException(422, "Nenhum departamento encontrado para o ID Orçamento e Cloud informados")

    def group(entries: list[PlanningEntry]) -> dict[int, dict[str, Any]]:
        grouped: dict[int, dict[str, Any]] = {}
        for entry in entries:
            if entry.department_id is None: continue
            grouped.setdefault(entry.department_id, {"entity_type": "department", "department_id": entry.department_id, "department_levels": entry.department_levels, "values": [0.0] * 12})["values"][entry.month - 1] = float(entry.value or 0)
        return grouped

    current_budget = (await db.execute(select(PlanningVersion).where(PlanningVersion.financial_type_id == budget_type.id, PlanningVersion.year == year, PlanningVersion.is_current == True))).scalars().first()
    budget_entries = (await db.execute(select(PlanningEntry).where(PlanningEntry.planning_version_id == current_budget.id))).scalars().all() if current_budget else []
    defense_entries = (await db.execute(select(PlanningEntry).where(PlanningEntry.planning_version_id == defense.id, PlanningEntry.department_id.in_(scope_ids)))).scalars().all()
    if not defense_entries:
        raise HTTPException(422, "A defesa não possui valores no escopo selecionado")
    merged = group(budget_entries)
    for department_id in scope_ids: merged.pop(department_id, None)
    merged.update(group(defense_entries))
    snapshot = await _publish_version(budget_type, year, list(merged.values()), "DEFESA_APROVADA", db, title=body.title, description=body.description, event_type="APROVACAO", effective_date=body.effective_date)
    snapshot.scope_id_orcamento = body.id_orcamento
    snapshot.scope_cloud = body.cloud_scope
    defense.scope_id_orcamento = body.id_orcamento
    defense.scope_cloud = body.cloud_scope
    defense.published_to_budget_at = datetime.now(timezone.utc)
    defense.published_budget_version_id = snapshot.id
    await db.commit()
    await db.refresh(snapshot)
    await db.refresh(defense)
    return {"budget_version": _out(snapshot), "defense_version": _out(defense), "rows": len(defense_entries)}


@router.post("/{type_code}/{year}/versions/{version_id}/restore")
async def restore_version(type_code: str, year: int, version_id: int, body: PlanningRestoreIn, db: AsyncSession = Depends(get_db)):
    fin_type = await _type_for_write(type_code, db)
    version = await db.get(PlanningVersion, version_id)
    if not version or version.financial_type_id != fin_type.id or version.year != year: raise HTTPException(404, "Snapshot não encontrado")
    entries = (await db.execute(select(PlanningEntry).where(PlanningEntry.planning_version_id == version.id, PlanningEntry.entity_type == "department"))).scalars().all()
    grouped: dict[tuple[str, int], dict[str, Any]] = {}
    for entry in entries:
        key = (entry.entity_type, entry.department_id)
        grouped.setdefault(key, {"entity_type": entry.entity_type, "department_id": entry.department_id, "department_levels": entry.department_levels, "values": [0.0] * 12})["values"][entry.month - 1] = entry.value
    snapshot = await _publish_version(fin_type, year, list(grouped.values()), "RESTORE", db, title=body.title, description=body.description, event_type=body.event_type, effective_date=body.effective_date, restored_from=version.id, scope_id_orcamento=body.scope_id_orcamento or version.scope_id_orcamento, scope_cloud=body.scope_cloud or version.scope_cloud)
    return {"version": _out(snapshot)}
